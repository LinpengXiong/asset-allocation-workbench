#!/usr/bin/env python3
"""Build a five-sheet, workbench-ready China macro scorecard proposal.

The detailed proposal and the user's original workbook remain untouched.
No current observations or investment conclusions are inserted.
"""

from __future__ import annotations

from pathlib import Path
from runpy import run_path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'deliverables' / '中国宏观评级评分框架_工作台精简版.xlsx'
SOURCE_RULES = run_path(str(Path(__file__).with_name('build-macro-score-workbook.py')))
CORE = SOURCE_RULES['CORE']
POLICY_METRICS = SOURCE_RULES['POLICY_METRICS']
POLICY_CORE = SOURCE_RULES['POLICY_CORE']

BURGUNDY = '72252A'
CREAM = 'F6F1EA'
PALE = 'FBF8F4'
INK = '282B2D'
MUTED = '62666A'
GOLD = 'B88B50'
INPUT = 'FFF4D9'
PILLAR_WEIGHT = {'经济运行': .30, '增长动力': .30, '金融信贷': .20, '宏观调控': .20}
POLICY_GROUP_WEIGHT = {'财政政策': .30, '货币政策': .30,
                       '产业政策': .20, '制度建设': .20}


def banner(ws, heading, subtitle, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, heading)
    ws.cell(1, 1).font = Font(name='PingFang SC', size=17, bold=True, color='FFFFFF')
    ws.cell(1, 1).fill = PatternFill('solid', fgColor=BURGUNDY)
    ws.cell(1, 1).alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 38
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(name='PingFang SC', size=10, color=MUTED)
    ws.cell(2, 1).fill = PatternFill('solid', fgColor=CREAM)
    ws.cell(2, 1).alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 33
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GOLD
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1


def header(ws, row, count):
    for cells in ws.iter_rows(min_row=row, max_row=row, max_col=count):
        for cell in cells:
            cell.font = Font(name='PingFang SC', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=BURGUNDY)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30


def body(ws, start, end, count):
    for r in range(start, end + 1):
        for row in ws.iter_rows(min_row=r, max_row=r, max_col=count):
            for cell in row:
                cell.font = Font(name='PingFang SC', size=10, color=INK)
                cell.fill = PatternFill('solid', fgColor='FFFFFF' if r % 2 else PALE)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.border = Border(bottom=Side(style='hair', color='E5DDD2'))
        ws.row_dimensions[r].height = 42


def widths(ws, mapping):
    for col, size in mapping.items():
        ws.column_dimensions[col].width = size


def bands(mode, limits, unit):
    if mode == '目标':
        a, b, c, d, e, f, g, h = limits
        return [f'<{a}或>{b} {unit}', f'{a}～{b} {unit}',
                f'{c}～{d} {unit}', f'{e}～{f} {unit}', f'{g}～{h} {unit}']
    a, b, c, d = limits
    if mode == '高优':
        return [f'<{a} {unit}', f'[{a},{b}) {unit}', f'[{b},{c}) {unit}',
                f'[{c},{d}) {unit}', f'≥{d} {unit}']
    return [f'≥{d} {unit}', f'[{c},{d}) {unit}', f'[{b},{c}) {unit}',
            f'[{a},{b}) {unit}', f'<{a} {unit}']


def numeric_formula(input_row, rule_row, mode):
    value = f'G{input_row}'
    rule = "'03_评分标准'!"
    if mode == '高优':
        return (f'IF({value}<{rule}O{rule_row},1,IF({value}<{rule}P{rule_row},2,'
                f'IF({value}<{rule}Q{rule_row},3,IF({value}<{rule}R{rule_row},4,5))))')
    if mode == '低优':
        return (f'IF({value}<{rule}O{rule_row},5,IF({value}<{rule}P{rule_row},4,'
                f'IF({value}<{rule}Q{rule_row},3,IF({value}<{rule}R{rule_row},2,1))))')
    expression = '1'
    for lo, hi, score in [('O','P',2), ('Q','R',3), ('S','T',4), ('U','V',5)]:
        expression = (f'IF(AND({value}>={rule}{lo}{rule_row},'
                      f'{value}<={rule}{hi}{rule_row}),{score},{expression})')
    return expression


def build():
    wb = Workbook()
    summary = wb.active
    summary.title = '01_评分总览'
    entry = wb.create_sheet('02_指标录入')
    rules = wb.create_sheet('03_评分标准')
    evidence = wb.create_sheet('04_政策原文依据')
    history = wb.create_sheet('05_历史记录')

    specs = []
    for item in CORE:
        specs.append(dict(id=item['id'], pillar=item['pillar'], group=item['group'],
                          name=item['name'], period=item['frequency'], unit=item['unit'],
                          mode=item['mode'], limits=item['limits'],
                          weight=PILLAR_WEIGHT[item['pillar']]*item['within'],
                          definition=item['definition']+'；'+item['note']))
    for ident,group,name,definition,period,unit,agency,mode,limits,part in POLICY_METRICS:
        specs.append(dict(id=ident, pillar='宏观调控', group=group, name=name,
                          period=period, unit=unit, mode=mode, limits=limits,
                          weight=.20*POLICY_GROUP_WEIGHT[group]*part,
                          definition=definition+'；首选原始机构：'+agency))
    for ident,group,name,within,agency,note,grades in POLICY_CORE[2:]:
        specs.append(dict(id=ident, pillar='宏观调控', group=group, name=name,
                          period='事件驱动／月度复核', unit='1—5分', mode='原文研判',
                          limits=[], weight=.20*within,
                          definition=note+'；首选原始机构：'+agency, grades=grades))
    assert len(specs) == 29
    assert abs(sum(item['weight'] for item in specs) - 1) < 1e-10

    banner(summary, '中国宏观评级｜评分总览',
           '四支柱固定权重30／30／20／20。政策分衡量逆周期动作与改革方向；经济结果由前三支柱验证。本期数值留空。', 7)
    summary['A4'] = '信息截止日'
    summary['B4'] = None
    summary['B4'].fill = PatternFill('solid', fgColor=INPUT)
    summary['B4'].number_format = 'yyyy-mm-dd'
    summary.append([])
    summary.append(['支柱','固定权重','已核验权重','覆盖率','得分贡献','已覆盖观察分','状态'])
    for r, pillar in enumerate(PILLAR_WEIGHT, 7):
        summary.cell(r, 1, pillar)
        summary.cell(r, 2, PILLAR_WEIGHT[pillar])
        summary.cell(r, 3, f'=SUMIFS(\'02_指标录入\'!$J$8:$J$36,\'02_指标录入\'!$B$8:$B$36,A{r},\'02_指标录入\'!$I$8:$I$36,">=1")')
        summary.cell(r, 4, f'=IFERROR(C{r}/B{r},0)')
        summary.cell(r, 5, f'=SUMIF(\'02_指标录入\'!$B$8:$B$36,A{r},\'02_指标录入\'!$K$8:$K$36)')
        summary.cell(r, 6, f'=IF(C{r}=0,"",E{r}/C{r})')
        summary.cell(r, 7, f'=IF(D{r}>=0.7,"可观察","不足")')
    summary['A12'] = '四支柱合计'
    for col in 'BCE':
        summary[f'{col}12'] = f'=SUM({col}7:{col}10)'
    summary['D12'] = '=IFERROR(C12/B12,0)'
    summary['F12'] = '=IF(C12=0,"",E12/C12)'
    summary['A14'] = '评级状态'
    summary['B14'] = '=IF(ROUND(D12,6)=1,"正式",IF(AND(D12>=0.85,COUNTIF(G7:G10,"可观察")=4),"暂定","资料不足"))'
    summary['A15'] = '正式总分'
    summary['B15'] = '=IF(B14="正式",E12,"")'
    summary['A16'] = '正式档位'
    summary['B16'] = '=IF(B15="","不定正式档",IF(B15>=4.2,"I档 强复苏",IF(B15>=3.4,"II档 温和改善",IF(B15>=2.6,"III档 中性平稳",IF(B15>=1.8,"IV档 偏弱承压","V档 明显下行")))))'
    summary['A18'] = '基本面观察分'
    summary['B18'] = '=IF(SUM(C7:C9)=0,"",SUM(E7:E9)/SUM(C7:C9))'
    summary['A19'] = '政策动作观察分'
    summary['B19'] = '=F10'
    summary['A21'] = '政策细分'
    summary['B21'] = '已覆盖观察分'
    summary['C21'] = '覆盖率'
    for r, group in enumerate(POLICY_GROUP_WEIGHT, 22):
        summary.cell(r, 1, group)
        summary.cell(r, 2, f'=IFERROR(SUMIF(\'02_指标录入\'!$C$8:$C$36,A{r},\'02_指标录入\'!$K$8:$K$36)/SUMIFS(\'02_指标录入\'!$J$8:$J$36,\'02_指标录入\'!$C$8:$C$36,A{r},\'02_指标录入\'!$I$8:$I$36,">=1"),"")')
        summary.cell(r, 3, f'=IFERROR(SUMIFS(\'02_指标录入\'!$J$8:$J$36,\'02_指标录入\'!$C$8:$C$36,A{r},\'02_指标录入\'!$I$8:$I$36,">=1")/(0.2*{POLICY_GROUP_WEIGHT[group]}),0)')
    summary.merge_cells('A28:G28')
    summary['A28'] = '缺失项不补3分；观察分只描述已覆盖数据，完整覆盖才产生可比的正式分数。阈值是研究试行标准，不是官方评级。'
    summary['A28'].alignment = Alignment(wrap_text=True, vertical='center')
    summary.row_dimensions[28].height = 38
    header(summary, 6, 7)
    header(summary, 21, 3)
    body(summary, 7, 10, 7)
    body(summary, 22, 25, 3)
    for r in range(7, 13):
        for c in (2, 3, 4): summary.cell(r, c).number_format = '0.0%'
        for c in (5, 6): summary.cell(r, c).number_format = '0.00'
    for r in range(22, 26): summary.cell(r, 3).number_format = '0.0%'
    for addr in ('B15','B18','B19'): summary[addr].number_format = '0.00'
    for r in (14,15,16,18,19): summary.cell(r, 1).font = Font(name='PingFang SC', bold=True, color=BURGUNDY)
    widths(summary, {'A':25,'B':22,'C':18,'D':16,'E':18,'F':21,'G':16})
    summary.freeze_panes = 'B7'

    banner(entry, '本期指标录入',
           '黄色格填统计期、发布日期、观测值、核验状态、官方链接与必要理由。原文评分从“政策原文依据”自动带入；不填本期数据就不生成评级。', 13)
    entry.merge_cells('A4:M4')
    entry['A4'] = '源顺序：官方原始发布 → 万得 Wind → 财新数据 API；后两级仅在前级无法取得可核验数据时使用，并记录代码、原始机构、期数、单位和口径。'
    entry['A4'].alignment = Alignment(wrap_text=True, vertical='center')
    entry.row_dimensions[4].height = 36
    entry.append([])
    entry.append([])
    entry.append(['ID','支柱','类别','核心指标','统计期','发布日期','值／原文分','核验','得分','固定权重','得分贡献','原文链接','口径／评分理由'])

    banner(rules, '五档评分标准',
           '一行一个核心指标；数字阈值和文本口径完整保留。财政、货币共六项量化规则；产业、制度沿用正式原文分档。隐藏的参数列只供自动计算。', 14)
    rules.append(['ID','支柱','类别','核心指标','频率','单位','规则方向','固定权重','1分','2分','3分','4分','5分','口径与边界'])
    header(rules, 3, 14)

    for i, spec in enumerate(specs):
        ir = 8 + i
        rr = 4 + i
        if spec['mode'] == '原文研判':
            grade_text = list(spec['grades'])
        else:
            grade_text = bands(spec['mode'], spec['limits'], spec['unit'])
        rules.append([spec['id'], spec['pillar'], spec['group'], spec['name'],
                      spec['period'], spec['unit'], spec['mode'], spec['weight'],
                      *grade_text, spec['definition']])
        if spec['mode'] != '原文研判':
            for j, threshold in enumerate(spec['limits'], 15):
                rules.cell(rr, j, threshold)
        entry.append([spec['id'], spec['pillar'], spec['group'], spec['name'],
                      None, None, None, '待核验', None,
                      f'=\'03_评分标准\'!H{rr}',
                      f'=IF(ISNUMBER(I{ir}),I{ir}*J{ir},0)', None, None])
        if spec['mode'] == '原文研判':
            er = 5 if spec['id'] == 'P03' else 6
            for col, src in [('E','B'),('F','C'),('G','H'),('H','I'),('L','D'),('M','G')]:
                entry[f'{col}{ir}'] = f'=\'04_政策原文依据\'!{src}{er}'
            extra = (f',\'04_政策原文依据\'!B{er}="",\'04_政策原文依据\'!C{er}="",'
                     f'\'04_政策原文依据\'!D{er}="",\'04_政策原文依据\'!E{er}="",'
                     f'\'04_政策原文依据\'!F{er}="",\'04_政策原文依据\'!G{er}="",'
                     f'G{ir}<1,G{ir}>5,MOD(G{ir},1)<>0')
            expression = f'G{ir}'
        else:
            extra = ''
            expression = numeric_formula(ir, rr, spec['mode'])
        entry[f'I{ir}'] = (f'=IF(OR(NOT(ISNUMBER(G{ir})),E{ir}="",'
                            f'NOT(ISNUMBER(F{ir})),F{ir}<=0,H{ir}<>"已核验",L{ir}="",'
                            f'NOT(ISNUMBER(\'01_评分总览\'!$B$4)),F{ir}>\'01_评分总览\'!$B$4'
                            f'{extra}),"",{expression})')
    body(entry, 8, 36, 13)
    body(rules, 4, 32, 14)
    for r in range(8, 37):
        entry.cell(r, 6).number_format = 'yyyy-mm-dd'
        entry.cell(r, 10).number_format = '0.00%'
        entry.cell(r, 11).number_format = '0.000'
        if r <= 34:
            for c in (5,6,7,8,12,13): entry.cell(r, c).fill = PatternFill('solid', fgColor=INPUT)
    for r in range(4, 33): rules.cell(r, 8).number_format = '0.00%'
    for col in 'OPQRSTUV': rules.column_dimensions[col].hidden = True
    status = DataValidation(type='list', formula1='"待核验,已核验,不可用"', allow_blank=False)
    entry.add_data_validation(status)
    status.add('H8:H34')
    entry.conditional_formatting.add('I8:I36', ColorScaleRule(
        start_type='num', start_value=1, start_color='F6DEDC',
        mid_type='num', mid_value=3, mid_color='FFF0CA',
        end_type='num', end_value=5, end_color='DDEEDB'))
    entry.auto_filter.ref = 'A7:M36'
    rules.auto_filter.ref = 'A3:N32'
    widths(entry, {'A':9,'B':16,'C':16,'D':47,'E':18,'F':18,'G':19,'H':15,'I':13,'J':15,'K':16,'L':54,'M':68})
    widths(rules, {'A':9,'B':16,'C':16,'D':45,'E':20,'F':15,'G':15,'H':15,
                   'I':31,'J':31,'K':31,'L':31,'M':31,'N':85})
    entry.freeze_panes = 'E8'
    rules.freeze_panes = 'E4'

    banner(evidence, '产业政策与制度建设｜原文评分依据',
           '只保留两个需要研判的政策类别。记录官方原文事实与研究判断，再按“评分标准”的五档规则填分；不等待增长或信贷效果兑现。', 10)
    evidence.append([])
    evidence.append(['类别','观察期','发布日期','官方原文链接','官方事实摘录','政策方向／覆盖范围','评分理由与主要约束','原文分（1—5）','核验状态','相关投资观察（待验证）'])
    evidence.append(['产业政策',None,None,None,None,None,None,None,'待核验',None])
    evidence.append(['制度建设',None,None,None,None,None,None,None,'待核验',None])
    header(evidence, 4, 10)
    body(evidence, 5, 6, 10)
    for r in (5,6):
        evidence.cell(r, 3).number_format = 'yyyy-mm-dd'
        for c in range(2,11): evidence.cell(r, c).fill = PatternFill('solid', fgColor=INPUT)
    grade = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True)
    evidence.add_data_validation(grade)
    grade.add('H5:H6')
    verify = DataValidation(type='list', formula1='"待核验,已核验,不可用"', allow_blank=False)
    evidence.add_data_validation(verify)
    verify.add('I5:I6')
    widths(evidence, {'A':18,'B':19,'C':18,'D':60,'E':65,'F':54,'G':65,'H':20,'I':17,'J':59})
    evidence.freeze_panes = 'D5'

    banner(history, '历史评级快照',
           '月度观察、季度定级后，将总览结果“粘贴为值”追加到新行；保留当时信息截止日与规则版本，不用修订后的数据覆盖旧评级。', 13)
    history.append(['类型','信息截止日','评价期','规则版本','经济运行','增长动力','金融信贷','宏观调控','正式总分','正式档位','覆盖率','政策依据／分歧','投研验证点'])
    for _ in range(24): history.append([None]*13)
    header(history, 3, 13)
    body(history, 4, 27, 13)
    history.auto_filter.ref = 'A3:M27'
    widths(history, {'A':18,'B':18,'C':18,'D':19,'E':16,'F':16,'G':16,'H':17,
                     'I':16,'J':20,'K':16,'L':67,'M':62})
    history.freeze_panes = 'E4'

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = 'auto'
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    return OUTPUT


if __name__ == '__main__':
    print(build())

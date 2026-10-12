#!/usr/bin/env python3
"""Build the versioned 29-indicator, four-pillar macro-rating scorecard.

The user-supplied compact workbook is read only. This is a rules proposal, so
no current observation or rating is filled in.
"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from runpy import run_path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('/Users/xionglinpeng/Desktop/中国宏观评级评分框架_工作台精简版.xlsx')
OUTPUT = ROOT / 'deliverables' / '中国宏观评级评分框架_V1.0试行定稿.xlsx'
STYLE = run_path(str(Path(__file__).with_name('build-macro-score-compact.py')))
banner, header, body, widths = (STYLE[name] for name in ('banner', 'header', 'body', 'widths'))
bands, numeric_formula = STYLE['bands'], STYLE['numeric_formula']

PILLARS = {'经济运行': .30, '增长动力': .30, '金融信贷': .20, '宏观调控': .20}
ORDER = [
    'E01', 'E02', 'E03', 'E04', 'E05', 'E06',
    'G01', 'G02', 'G03', 'G05', 'G07', 'G08', 'G13', 'G09', 'G12', 'G10', 'G11',
    'F01', 'F02', 'F03', 'F05',
    'P01A', 'P01B', 'P01C', 'P02A', 'P02B', 'P02C', 'P03', 'P04',
]
WEIGHTS = {
    'E01': .085, 'E02': .04, 'E03': .03, 'E04': .015, 'E05': .065, 'E06': .065,
    'G01': .03, 'G02': .02, 'G03': .03, 'G05': .02,
    'G07': .025, 'G08': .025, 'G13': .025,
    'G09': .045, 'G12': .03, 'G10': .025, 'G11': .025,
    'F01': .04, 'F02': .05, 'F03': .07, 'F05': .04,
    'P01A': .03, 'P01B': .02, 'P01C': .01,
    'P02A': .03, 'P02B': .02, 'P02C': .01,
    'P03': .04, 'P04': .04,
}
OBSERVATIONS = [
    ('G04', '房屋新开工面积增速', '保留房地产供给侧观察，核心评分用销售与到位资金分别代表需求和融资。'),
    ('G06', '固定资产投资总额增速', '总量与制造业、民间、基建分项重叠；保留作分项一致性检查。'),
    ('F04', '人民币贷款存量增速', '与社融存量有较多共同信息但口径不同；保留作信用结构验证。'),
]
NEW = {
    'G11': dict(pillar='增长动力', group='前瞻需求', name='制造业PMI新订单指数',
                frequency='月度', unit='指数点', mode='高优', limits=[47, 49, 50, 52],
                definition='国家统计局制造业PMI新订单分项，经季节调整；扩散指数反映需求方向，并非订单量同比。五档为试行阈值。'),
    'G12': dict(pillar='增长动力', group='消费', name='服务零售额增速（年初累计同比，%）',
                frequency='月度', unit='%', mode='高优', limits=[0, 3, 5, 8],
                definition='国家统计局服务零售额年初累计同比；与社会消费品零售总额的当月同比频率及覆盖范围不同，不拼为一条线。五档为试行阈值。'),
    'G13': dict(pillar='增长动力', group='投资', name='基础设施投资增速（年初累计同比，%）',
                frequency='月度', unit='%', mode='高优', limits=[-5, 0, 3, 7],
                definition='国家统计局固定资产投资（不含农户）中的基础设施投资，采用当期官方定义（含电力、热力、燃气及水生产和供应业）；年初累计、名义增速。五档为试行阈值。'),
}

WEIGHT_RATIONALE = {
    'E01': '总量增长锚，覆盖面广；季度更新较慢。',
    'E02': '居民总体价格信号；与核心CPI共享信息。',
    'E03': '补充内需价格；与总体CPI共享信息。',
    'E04': '工业价格和利润线索；成本冲击及波动较大。',
    'E05': '就业是宏观运行的直接结果。',
    'E06': '较及时的景气信号；与产出和新订单部分重叠。',
    'G01': '实体生产主干；与利润及PMI部分重叠。',
    'G02': '补充企业盈利；波动和发布时间滞后。',
    'G03': '房地产需求主指标；与居民长贷联动。',
    'G05': '房企融资约束补充；与信用指标部分重叠。',
    'G07': '制造业资本开支和扩产意愿。',
    'G08': '民间投资与企业信心；与制造业投资部分重叠。',
    'G13': '政策支持下的实际投资活动，不重复算政策动作。',
    'G09': '商品消费需求主指标，但不覆盖全部服务。',
    'G12': '补充服务消费；与社零频率和范围不同。',
    'G10': '外需线索；美元计价受价格及汇率扰动。',
    'G11': '较早的需求变化线索；权重偏低待检验。',
    'F01': '居民信用结构；与房屋销售联动。',
    'F02': '企业中长期融资；与社融及投资部分重叠。',
    'F03': '信用总量锚；与分部门贷款重叠风险待检验。',
    'F05': '货币活化补充；新旧M1口径须分段。',
    'P01A': '正式财政资源安排，评价政策动作而非效果。',
    'P01B': '预算执行节奏；受季节安排影响。',
    'P01C': '专项债发行进度；前移不等于新增力度。',
    'P02A': '货币政策价格工具的直接动作。',
    'P02B': '货币政策数量工具；须固定机构口径。',
    'P02C': '窄口径定向工具，控制其总分影响。',
    'P03': '保留产业政策及供给侧信号；人工评分需原文复核。',
    'P04': '保留制度变动前瞻信息；人工评分需原文复核。',
}


def display_bands(mode, limits, unit):
    """Show the nested target rule as five disjoint, formula-equivalent bands."""
    if mode != '目标':
        return bands(mode, limits, unit)
    a, b, c, d, e, f, g, h = limits
    assert a < c < e < g <= h < f < d < b
    return [f'<{a}或>{b} {unit}',
            f'[{a},{c})或({d},{b}] {unit}',
            f'[{c},{e})或({f},{d}] {unit}',
            f'[{e},{g})或({h},{f}] {unit}',
            f'[{g},{h}] {unit}']


def source_specs():
    workbook = load_workbook(SOURCE, read_only=True, data_only=False)
    sheet = workbook['03_评分标准']
    specs = {}
    for row in range(4, 33):
        ident = sheet.cell(row, 1).value
        if not ident:
            continue
        specs[ident] = dict(
            pillar=sheet.cell(row, 2).value,
            group=sheet.cell(row, 3).value,
            name=sheet.cell(row, 4).value,
            frequency=sheet.cell(row, 5).value,
            unit=sheet.cell(row, 6).value,
            mode=sheet.cell(row, 7).value,
            grades=[sheet.cell(row, col).value for col in range(9, 14)],
            definition=sheet.cell(row, 14).value,
            limits=[sheet.cell(row, col).value for col in range(15, 23)
                    if sheet.cell(row, col).value is not None],
        )
    specs.update(NEW)
    specs['P02C']['name'] = '其他结构性货币政策工具近3月净投放/GDP'
    specs['P02C']['definition'] = (
        '仅指人民银行月度“中央银行各项工具流动性投放情况”中的“其他结构性货币政策工具”净投放，'
        '不是全部结构性工具；近3月合计除以最近四季名义GDP。官方该口径缺月或变动时留空。五档为试行阈值。'
    )
    specs['P03']['grades'] = [
        '正式措施明显损害有效竞争或供给预期，且缺少纠偏安排',
        '方向或执行规则不清，市场预期不确定性显著上升',
        '经核查无实质新动作、既有政策延续；或正负作用大体抵消',
        '正式文件明确供给侧治理或产业支持工具、范围与时点',
        '跨部门工具、资金或标准及试点时间表明确；不等待订单兑现',
    ]
    specs['P04']['grades'] = [
        '正式制度变化明显损害准入、产权或规则稳定性',
        '正式方向偏限制或关键执行规则不清，不确定性上升',
        '经核查无实质新动作、既有规则延续；或正负作用大体抵消',
        '正式法律或制度文本明确适用范围、实施办法和生效时点',
        '关键制度改革有明确规则、责任机构和时间表；不等待盈利兑现',
    ]
    for ident in ORDER:
        specs[ident]['definition'] += f"；权重依据：{WEIGHT_RATIONALE[ident]}"
    assert len(ORDER) == len(set(ORDER)) == len(WEIGHTS) == len(WEIGHT_RATIONALE) == 29
    assert set(ORDER).issubset(specs)
    totals = defaultdict(float)
    for ident in ORDER:
        totals[specs[ident]['pillar']] += WEIGHTS[ident]
    for pillar, weight in PILLARS.items():
        assert abs(totals[pillar] - weight) < 1e-10, (pillar, totals[pillar])
    return specs


def build():
    specs = source_specs()
    wb = Workbook()
    summary = wb.active
    summary.title = '01_评分总览'
    entry = wb.create_sheet('02_指标录入')
    rules = wb.create_sheet('03_评分标准')
    history = wb.create_sheet('04_历史记录')
    last_entry = 7 + len(ORDER)
    last_rule = 3 + len(ORDER)

    banner(summary, '中国宏观评级｜V1.0试行定稿',
           '固定试行规则：经济运行30%、增长动力30%、金融信贷20%、宏观调控20%；29项核心评分，3项辅助观察。本期未填值，不生成评级。', 7)
    summary['A4'] = '信息截止日'
    summary['B4'].fill = PatternFill('solid', fgColor=STYLE['INPUT'])
    summary['B4'].number_format = 'yyyy-mm-dd'
    summary.append([])
    summary.append(['支柱', '固定权重', '已核验权重', '覆盖率', '得分贡献', '已覆盖观察分', '状态'])
    for row, (pillar, weight) in enumerate(PILLARS.items(), 7):
        summary.cell(row, 1, pillar)
        summary.cell(row, 2, weight)
        summary.cell(row, 3, f'=SUMIFS(\'02_指标录入\'!$J$8:$J${last_entry},\'02_指标录入\'!$B$8:$B${last_entry},A{row},\'02_指标录入\'!$I$8:$I${last_entry},">=1")')
        summary.cell(row, 4, f'=IFERROR(C{row}/B{row},0)')
        summary.cell(row, 5, f'=SUMIF(\'02_指标录入\'!$B$8:$B${last_entry},A{row},\'02_指标录入\'!$K$8:$K${last_entry})')
        summary.cell(row, 6, f'=IF(C{row}=0,"",E{row}/C{row})')
        summary.cell(row, 7, f'=IF(D{row}>=0.7,"可观察","不足")')
    summary['A12'] = '四支柱合计'
    for col in 'BCE':
        summary[f'{col}12'] = f'=SUM({col}7:{col}10)'
    summary['D12'] = '=IFERROR(C12/B12,0)'
    summary['F12'] = '=IF(C12=0,"",E12/C12)'
    summary['A14'] = '评级状态'
    summary['B14'] = '=IF(ROUND(D12,6)=1,"正式",IF(AND(D12>=0.85,COUNTIF(G7:G10,"可观察")=4),"暂定","资料不足"))'
    summary['A15'] = '正式总分'
    summary['B15'] = '=IF(B14="正式",E12,"")'
    summary['A16'] = '正式档位'
    summary['B16'] = '=IF(B15="","不定正式档",IF(B15>=4.2,"I档 强复苏",IF(B15>=3.4,"II档 温和改善",IF(B15>=2.6,"III档 中性平稳",IF(B15>=1.8,"IV档 偏弱承压","V档 明显下行")))))'
    summary['A18'] = '基本面观察分'
    summary['B18'] = '=IF(SUM(C7:C9)=0,"",SUM(E7:E9)/SUM(C7:C9))'
    summary['A19'] = '政策动作观察分'
    summary['B19'] = '=F10'
    summary.append([])
    summary['A21'], summary['B21'], summary['C21'] = '政策细分', '已覆盖观察分', '覆盖率'
    for row, group in enumerate(['财政政策', '货币政策', '产业政策', '制度建设'], 22):
        summary.cell(row, 1, group)
        selected = f'SUMIFS(\'02_指标录入\'!$J$8:$J${last_entry},\'02_指标录入\'!$C$8:$C${last_entry},A{row},\'02_指标录入\'!$I$8:$I${last_entry},">=1")'
        fixed = f'SUMIF(\'02_指标录入\'!$C$8:$C${last_entry},A{row},\'02_指标录入\'!$J$8:$J${last_entry})'
        contribution = f'SUMIF(\'02_指标录入\'!$C$8:$C${last_entry},A{row},\'02_指标录入\'!$K$8:$K${last_entry})'
        summary.cell(row, 2, f'=IFERROR({contribution}/{selected},"")')
        summary.cell(row, 3, f'=IFERROR({selected}/{fixed},0)')
    summary.merge_cells('A28:G28')
    summary['A28'] = '未核验项不补3分；覆盖不足不发布正式评级。政策动作不要求经济效果当期兑现，效果由前三支柱验证。五档阈值仍需历史回测。'
    summary['A28'].alignment = Alignment(wrap_text=True, vertical='center')
    summary.row_dimensions[28].height = 38
    summary['A30'], summary['B30'] = '规则版本', 'V1.0（2026-10-04）'
    summary['A31'], summary['B31'] = '验证状态', '可按同一规则试行；细项权重与阈值尚待历史回测，非统计最优。'
    summary.merge_cells('B31:G31')
    summary['B31'].alignment = Alignment(wrap_text=True, vertical='center')
    summary.row_dimensions[31].height = 32
    header(summary, 6, 7)
    header(summary, 21, 3)
    body(summary, 7, 10, 7)
    body(summary, 22, 25, 3)
    for row in range(7, 13):
        for col in (2, 3, 4):
            summary.cell(row, col).number_format = '0.0%'
        for col in (5, 6):
            summary.cell(row, col).number_format = '0.00'
    for row in range(22, 26):
        summary.cell(row, 3).number_format = '0.0%'
    for address in ('B15', 'B18', 'B19'):
        summary[address].number_format = '0.00'
    widths(summary, {'A': 25, 'B': 22, 'C': 18, 'D': 16, 'E': 18, 'F': 21, 'G': 16})
    summary.freeze_panes = 'B7'

    banner(entry, '本期指标录入｜V1.0试行',
           '黄色格填写统计期、发布日期、观测值、核验状态、官方原文和理由。产业与制度两项直接按正式原文人工给1—5分，不再引用已删除的工作表。', 13)
    entry.merge_cells('A4:M4')
    entry['A4'] = '数据来源依次为官方原始发布 → Wind → 财新数据API；后级仅在前级取不到所需可核验数据时使用。不同频率及新旧统计口径不得拼接。'
    entry['A4'].alignment = Alignment(wrap_text=True, vertical='center')
    entry.row_dimensions[4].height = 34
    entry.append([])
    entry.append([])
    entry.append(['ID', '支柱', '类别', '核心指标', '统计期', '发布日期', '值／原文分', '核验', '得分', '固定权重', '得分贡献', '原文链接', '口径／评分理由'])
    header(entry, 7, 13)

    banner(rules, '核心评分标准｜V1.0试行',
           '29项固定权重合计100%；四个目标区间指标的五档文字已改成互斥区间。财政、货币各三项量化指标；产业、制度按正式原文五档评分。', 14)
    rules.append(['ID', '支柱', '类别', '核心指标', '频率', '单位', '规则方向', '总分权重', '1分', '2分', '3分', '4分', '5分', '口径与边界'])
    header(rules, 3, 14)

    for offset, ident in enumerate(ORDER):
        spec = specs[ident]
        input_row, rule_row = 8 + offset, 4 + offset
        grades = spec.get('grades') if spec['mode'] == '原文研判' else display_bands(spec['mode'], spec['limits'], spec['unit'])
        rules.append([ident, spec['pillar'], spec['group'], spec['name'], spec['frequency'],
                      spec['unit'], spec['mode'], WEIGHTS[ident], *grades, spec['definition']])
        if spec['mode'] != '原文研判':
            for col, limit in enumerate(spec['limits'], 15):
                rules.cell(rule_row, col, limit)
        entry.append([ident, spec['pillar'], spec['group'], spec['name'], None, None, None,
                      '待核验', None, f'=\'03_评分标准\'!H{rule_row}',
                      f'=IF(ISNUMBER(I{input_row}),I{input_row}*J{input_row},0)', None, None])
        value = f'G{input_row}'
        extra = f',M{input_row}="",{value}<1,{value}>5,MOD({value},1)<>0' if spec['mode'] == '原文研判' else ''
        expression = value if spec['mode'] == '原文研判' else numeric_formula(input_row, rule_row, spec['mode'])
        entry[f'I{input_row}'] = (
            f'=IF(OR(NOT(ISNUMBER({value})),E{input_row}="",NOT(ISNUMBER(F{input_row})),'
            f'F{input_row}<=0,H{input_row}<>"已核验",L{input_row}="",'
            f'NOT(ISNUMBER(\'01_评分总览\'!$B$4)),F{input_row}>\'01_评分总览\'!$B$4'
            f'{extra}),"",{expression})'
        )
    body(entry, 8, last_entry, 13)
    body(rules, 4, last_rule, 14)
    for row in range(8, last_entry + 1):
        entry.cell(row, 6).number_format = 'yyyy-mm-dd'
        entry.cell(row, 10).number_format = '0.00%'
        entry.cell(row, 11).number_format = '0.000'
        for col in (5, 6, 7, 8, 12, 13):
            entry.cell(row, col).fill = PatternFill('solid', fgColor=STYLE['INPUT'])
    for row in range(4, last_rule + 1):
        rules.cell(row, 8).number_format = '0.00%'
    for col in 'OPQRSTUV':
        rules.column_dimensions[col].hidden = True
    status = DataValidation(type='list', formula1='"待核验,已核验,不可用"', allow_blank=False)
    entry.add_data_validation(status)
    status.add(f'H8:H{last_entry}')
    grade = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True)
    entry.add_data_validation(grade)
    for ident in ('P03', 'P04'):
        grade.add(f'G{8 + ORDER.index(ident)}')
    entry.conditional_formatting.add(f'I8:I{last_entry}', ColorScaleRule(
        start_type='num', start_value=1, start_color='F6DEDC',
        mid_type='num', mid_value=3, mid_color='FFF0CA',
        end_type='num', end_value=5, end_color='DDEEDB'))
    entry.auto_filter.ref = f'A7:M{last_entry}'
    rules.auto_filter.ref = f'A3:N{last_rule}'
    widths(entry, {'A': 9, 'B': 16, 'C': 16, 'D': 47, 'E': 18, 'F': 18,
                   'G': 19, 'H': 15, 'I': 13, 'J': 15, 'K': 16, 'L': 54, 'M': 68})
    widths(rules, {'A': 9, 'B': 16, 'C': 16, 'D': 45, 'E': 21, 'F': 15,
                   'G': 15, 'H': 15, 'I': 31, 'J': 31, 'K': 31, 'L': 31, 'M': 31, 'N': 85})
    entry.freeze_panes = 'E8'
    rules.freeze_panes = 'E4'

    observation_entry_title = last_entry + 3
    observation_entry_header = observation_entry_title + 1
    entry.merge_cells(start_row=observation_entry_title, start_column=1,
                      end_row=observation_entry_title, end_column=13)
    entry.cell(observation_entry_title, 1, '辅助观察｜照常逐期录入，但不计入评级和覆盖率')
    entry.cell(observation_entry_title, 1).font = Font(
        name='PingFang SC', bold=True, color=STYLE['BURGUNDY'])
    for col, label in enumerate(['ID', '支柱', '类别', '观察指标', '统计期', '发布日期',
                                 '观测值', '核验', '不计分', '权重', '得分贡献',
                                 '原文链接', '口径／观察用途'], 1):
        entry.cell(observation_entry_header, col, label)
    header(entry, observation_entry_header, 13)
    for row, (ident, name, reason) in enumerate(OBSERVATIONS,
                                               observation_entry_header + 1):
        spec = specs[ident]
        for col, value in {1: ident, 2: spec['pillar'], 3: spec['group'],
                           4: name, 8: '待核验', 9: '—', 10: 0,
                           13: reason}.items():
            entry.cell(row, col, value)
        entry.cell(row, 6).number_format = 'yyyy-mm-dd'
        entry.cell(row, 10).number_format = '0.00%'
        for col in (5, 6, 7, 8, 12, 13):
            entry.cell(row, col).fill = PatternFill('solid', fgColor=STYLE['INPUT'])
    observation_entry_end = observation_entry_header + len(OBSERVATIONS)
    body(entry, observation_entry_header + 1, observation_entry_end, 13)
    for row in range(observation_entry_header + 1, observation_entry_end + 1):
        for col in (5, 6, 7, 8, 12, 13):
            entry.cell(row, col).fill = PatternFill('solid', fgColor=STYLE['INPUT'])
    status.add(f'H{observation_entry_header + 1}:H{observation_entry_end}')

    observation_title = last_rule + 3
    rules.merge_cells(start_row=observation_title, start_column=1, end_row=observation_title, end_column=14)
    rules.cell(observation_title, 1, '辅助观察｜不参与总分（原指标保留，权重为0）')
    rules.cell(observation_title, 1).font = Font(name='PingFang SC', bold=True, color=STYLE['BURGUNDY'])
    for row, (ident, name, reason) in enumerate(OBSERVATIONS, observation_title + 1):
        rules.cell(row, 1, ident)
        rules.cell(row, 4, name)
        rules.cell(row, 8, 0)
        rules.cell(row, 14, reason)
        rules.cell(row, 8).number_format = '0.00%'
        rules.cell(row, 14).alignment = Alignment(wrap_text=True, vertical='center')
        rules.row_dimensions[row].height = 40

    banner(history, '历史评级快照',
           '每次定级后仅粘贴当时已核验的值与规则版本（当前为V1.0）；不以今后修订值覆盖旧评级。新旧版本不能直接比较分数变化。', 13)
    history.append(['类型', '信息截止日', '评价期', '规则版本', '经济运行', '增长动力', '金融信贷',
                    '宏观调控', '正式总分', '正式档位', '覆盖率', '政策依据／分歧', '投研验证点'])
    for _ in range(24):
        history.append([None] * 13)
    header(history, 3, 13)
    body(history, 4, 27, 13)
    history.auto_filter.ref = 'A3:M27'
    widths(history, {'A': 18, 'B': 18, 'C': 18, 'D': 19, 'E': 16, 'F': 16,
                     'G': 16, 'H': 17, 'I': 16, 'J': 20, 'K': 16, 'L': 67, 'M': 62})
    history.freeze_panes = 'E4'

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = 'auto'
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    return OUTPUT


if __name__ == '__main__':
    print(build())

#!/usr/bin/env python3
"""Export the approved trial scorecard for the workbench without inventing observations."""

import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables/中国宏观评级评分框架_V1.0试行定稿.xlsx"
TARGET = ROOT / "app/data/macro-rating-v1.json"
PILLARS = {"经济运行": 0.30, "增长动力": 0.30, "金融信贷": 0.20, "宏观调控": 0.20}


def main():
    # This exporter belongs to the archived V1.0 workbook, not the live rules.
    if TARGET.exists():
        current_version = json.loads(TARGET.read_text(encoding="utf-8")).get("version")
        if current_version != "V1.0":
            raise SystemExit(
                f"Refusing to overwrite the live {current_version} rules with the archived V1.0 workbook. "
                "Export a historical workbook to a separate file instead."
            )
    book = load_workbook(SOURCE, read_only=True, data_only=False)
    rules = book["03_评分标准"]
    entries = book["02_指标录入"]
    indicators = []
    totals = defaultdict(float)
    for row in range(4, 33):
        ident = rules.cell(row, 1).value
        if not ident:
            continue
        entry_row = row + 4
        assert entries.cell(entry_row, 1).value == ident
        assert entries.cell(entry_row, 7).value is None, f"Current value must remain blank: {ident}"
        item = {
            "id": ident,
            "pillar": rules.cell(row, 2).value,
            "group": rules.cell(row, 3).value,
            "name": rules.cell(row, 4).value,
            "frequency": rules.cell(row, 5).value,
            "unit": rules.cell(row, 6).value,
            "mode": rules.cell(row, 7).value,
            "weight": rules.cell(row, 8).value,
            "grades": [rules.cell(row, col).value for col in range(9, 14)],
            "definition": rules.cell(row, 14).value,
            "limits": [rules.cell(row, col).value for col in range(15, 23)
                       if rules.cell(row, col).value is not None],
        }
        assert all(isinstance(grade, str) and grade for grade in item["grades"])
        indicators.append(item)
        totals[item["pillar"]] += item["weight"]
    assert len(indicators) == 29 and len({item["id"] for item in indicators}) == 29
    assert all(abs(totals[name] - weight) < 1e-9 for name, weight in PILLARS.items())
    auxiliary = []
    for row in range(41, 44):
        ident = entries.cell(row, 1).value
        assert ident in {"G04", "G06", "F04"}, (row, ident)
        assert entries.cell(row, 7).value is None
        auxiliary.append({
            "id": ident,
            "pillar": entries.cell(row, 2).value,
            "group": entries.cell(row, 3).value,
            "name": entries.cell(row, 4).value,
            "reason": entries.cell(row, 13).value,
            "weight": 0,
        })
    payload = {
        "version": "V1.0",
        "versionDate": "2026-10-04",
        "status": "试行定稿，权重和阈值尚待历史回测",
        "pillars": PILLARS,
        "ratingBands": [
            {"minimum": 4.2, "label": "I档 强复苏"},
            {"minimum": 3.4, "label": "II档 温和改善"},
            {"minimum": 2.6, "label": "III档 中性平稳"},
            {"minimum": 1.8, "label": "IV档 偏弱承压"},
            {"minimum": 1.0, "label": "V档 明显下行"},
        ],
        "coverage": {"formal": 1.0, "provisional": 0.85, "eachPillar": 0.70},
        "indicators": indicators,
        "auxiliary": auxiliary,
    }
    TARGET.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Exported {len(indicators)} core rules and {len(auxiliary)} auxiliary rows to {TARGET}")


if __name__ == "__main__":
    main()
